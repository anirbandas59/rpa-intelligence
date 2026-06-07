"""
File storage service for bulk upload sessions.

Manages uploaded file lifecycle: storage, parsing, cleanup. Files stored in
backend/data/temp/uploads/ with UUID names. Supports CSV and XLSX formats.
Memory-efficient chunked parsing for large files (yields batches of 1000 rows).
"""

import csv
import uuid
from collections.abc import AsyncIterator
from datetime import datetime, timedelta
from pathlib import Path

from core.exceptions import DocumentProcessingError
from db.models import UploadSession, UseCase

try:
    import openpyxl
except ImportError:
    openpyxl = None


class FileStorageService:
    """Handles file storage and parsing for bulk upload sessions."""

    def __init__(self, upload_dir: str | None = None):
        """Initialize with upload directory path."""
        if upload_dir is None:
            # Default to backend/data/temp/uploads
            base_dir = Path(__file__).parent.parent
            upload_dir = str(base_dir / "data" / "temp" / "uploads")

        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def create_upload_session(
        self,
        file_content: bytes,
        filename: str,
        project_id: str,
        user_id: str,
    ) -> UploadSession:
        """
        Store uploaded file and create session record.

        Args:
            file_content: Raw file bytes
            filename: Original filename
            project_id: Parent project UUID
            user_id: Uploader UUID

        Returns:
            UploadSession with preview data

        Raises:
            DocumentProcessingError: If file parsing fails
        """
        if not file_content:
            raise DocumentProcessingError("File is empty")

        # Generate UUID for stored file
        session_id = str(uuid.uuid4())
        file_ext = Path(filename).suffix
        stored_filename = f"{session_id}{file_ext}"
        file_path = self.upload_dir / stored_filename

        # Write file to disk
        file_path.write_bytes(file_content)

        # Parse columns and count rows
        try:
            columns, row_count = self._parse_file_metadata(file_path, filename)
        except Exception as e:
            # Clean up file if parsing fails
            file_path.unlink(missing_ok=True)
            raise DocumentProcessingError(f"Failed to parse file: {str(e)}") from e

        # Create session record
        session = UploadSession(
            id=session_id,
            project_id=project_id,
            user_id=user_id,
            filename=filename,
            file_path=str(file_path),
            file_size=len(file_content),
            columns={"columns": columns},  # JSON array wrapper
            row_count=row_count,
            status="preview",
            created_at=datetime.utcnow(),
            expires_at=datetime.utcnow() + timedelta(hours=24),
        )

        return session

    def get_preview_rows(self, session: UploadSession, limit: int = 5) -> list[dict[str, str]]:
        """
        Read first N rows from stored file.

        Args:
            session: UploadSession with file_path
            limit: Number of rows to return

        Returns:
            List of row dicts (column -> value)

        Raises:
            DocumentProcessingError: If file read fails
        """
        file_path = Path(session.file_path)
        if not file_path.exists():
            raise DocumentProcessingError("Uploaded file no longer exists")

        try:
            if session.filename.endswith(".csv"):
                return self._preview_csv(file_path, limit)
            elif session.filename.endswith(".xlsx"):
                return self._preview_xlsx(file_path, limit)
            else:
                raise DocumentProcessingError(f"Unsupported file type: {session.filename}")
        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to read preview: {str(e)}") from e

    async def parse_file_with_mapping(
        self,
        session: UploadSession,
        column_mapping: dict[str, str | None],
        batch_size: int = 1000,
    ) -> AsyncIterator[list[UseCase]]:
        """
        Generator yielding UseCase objects in batches.

        Memory-efficient for large files. Reads file in chunks, applies
        column mapping, yields UseCase ORM objects ready for bulk insert.

        Args:
            session: UploadSession with file_path and project_id
            column_mapping: {name, description, source_platform, install_status}
            batch_size: Rows per batch (default 1000)

        Yields:
            Lists of UseCase objects

        Raises:
            DocumentProcessingError: If file parsing fails
        """
        file_path = Path(session.file_path)
        if not file_path.exists():
            raise DocumentProcessingError("Uploaded file no longer exists")

        try:
            if session.filename.endswith(".csv"):
                async for batch in self._parse_csv_batches(
                    file_path, session.project_id, column_mapping, batch_size
                ):
                    yield batch
            elif session.filename.endswith(".xlsx"):
                async for batch in self._parse_xlsx_batches(
                    file_path, session.project_id, column_mapping, batch_size
                ):
                    yield batch
            else:
                raise DocumentProcessingError(f"Unsupported file type: {session.filename}")
        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to parse file: {str(e)}") from e

    def cleanup_expired_sessions(self, sessions: list[UploadSession]) -> int:
        """
        Delete files for expired sessions.

        Args:
            sessions: List of UploadSession records to clean up

        Returns:
            Number of files deleted
        """
        deleted = 0
        for session in sessions:
            file_path = Path(session.file_path)
            if file_path.exists():
                file_path.unlink()
                deleted += 1
        return deleted

    def _parse_file_metadata(self, file_path: Path, filename: str) -> tuple[list[str], int]:
        """
        Extract column names and row count.

        Args:
            file_path: Path to stored file
            filename: Original filename (for type detection)

        Returns:
            (columns, row_count)

        Raises:
            DocumentProcessingError: If parsing fails
        """
        if filename.endswith(".csv"):
            return self._parse_csv_metadata(file_path)
        elif filename.endswith(".xlsx"):
            return self._parse_xlsx_metadata(file_path)
        else:
            raise DocumentProcessingError(
                f"Unsupported file type: {filename}. Only .csv and .xlsx are supported."
            )

    def _parse_csv_metadata(self, file_path: Path) -> tuple[list[str], int]:
        """Parse CSV column names and count rows."""
        try:
            with file_path.open("r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                if not reader.fieldnames:
                    raise DocumentProcessingError("CSV has no columns")
                columns = list(reader.fieldnames)
                row_count = sum(1 for _ in reader)
            return columns, row_count
        except UnicodeDecodeError as e:
            raise DocumentProcessingError(
                "File encoding error. Please ensure the file is UTF-8 encoded."
            ) from e

    def _parse_xlsx_metadata(self, file_path: Path) -> tuple[list[str], int]:
        """Parse XLSX column names and count rows."""
        if openpyxl is None:
            raise DocumentProcessingError("openpyxl is not installed")

        workbook = openpyxl.load_workbook(file_path, read_only=True)
        ws = workbook.active

        # Get columns from first row
        columns = []
        for cell in ws[1]:
            if cell.value:
                columns.append(str(cell.value))

        if not columns:
            raise DocumentProcessingError("XLSX has no columns")

        # Count rows (exclude header)
        row_count = ws.max_row - 1 if ws.max_row > 1 else 0

        workbook.close()
        return columns, row_count

    def _preview_csv(self, file_path: Path, limit: int) -> list[dict[str, str]]:
        """Read first N rows from CSV."""
        preview = []
        with file_path.open("r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for i, row in enumerate(reader):
                if i >= limit:
                    break
                preview.append(dict(row))
        return preview

    def _preview_xlsx(self, file_path: Path, limit: int) -> list[dict[str, str]]:
        """Read first N rows from XLSX."""
        if openpyxl is None:
            raise DocumentProcessingError("openpyxl is not installed")

        workbook = openpyxl.load_workbook(file_path, read_only=True)
        ws = workbook.active

        # Get columns
        columns = []
        for cell in ws[1]:
            if cell.value:
                columns.append(str(cell.value))

        # Get preview rows
        preview = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            if i >= limit:
                break
            row_dict = {
                columns[j]: str(val) if val is not None else ""
                for j, val in enumerate(row)
                if j < len(columns)
            }
            preview.append(row_dict)

        workbook.close()
        return preview

    async def _parse_csv_batches(
        self,
        file_path: Path,
        project_id: str,
        column_mapping: dict[str, str | None],
        batch_size: int,
    ) -> AsyncIterator[list[UseCase]]:
        """Parse CSV in batches, yielding UseCase objects."""
        with file_path.open("r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            batch = []

            for row in reader:
                use_case = self._row_to_use_case(row, project_id, column_mapping)
                if use_case:  # Skip if name is empty
                    batch.append(use_case)

                if len(batch) >= batch_size:
                    yield batch
                    batch = []

            # Yield remaining
            if batch:
                yield batch

    async def _parse_xlsx_batches(
        self,
        file_path: Path,
        project_id: str,
        column_mapping: dict[str, str | None],
        batch_size: int,
    ) -> AsyncIterator[list[UseCase]]:
        """Parse XLSX in batches, yielding UseCase objects."""
        if openpyxl is None:
            raise DocumentProcessingError("openpyxl is not installed")

        workbook = openpyxl.load_workbook(file_path, read_only=True)
        ws = workbook.active

        # Get columns
        columns = []
        for cell in ws[1]:
            if cell.value:
                columns.append(str(cell.value))

        batch = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {
                columns[j]: str(val) if val is not None else ""
                for j, val in enumerate(row)
                if j < len(columns)
            }
            use_case = self._row_to_use_case(row_dict, project_id, column_mapping)
            if use_case:
                batch.append(use_case)

            if len(batch) >= batch_size:
                yield batch
                batch = []

        # Yield remaining
        if batch:
            yield batch

        workbook.close()

    def _row_to_use_case(
        self,
        row: dict[str, str],
        project_id: str,
        column_mapping: dict[str, str | None],
    ) -> UseCase | None:
        """
        Convert row dict to UseCase object.

        Args:
            row: Column -> value mapping
            project_id: Parent project UUID
            column_mapping: Field -> column name mapping

        Returns:
            UseCase object or None if name is empty
        """
        # Get name from mapped column
        name_col = column_mapping.get("name")
        if not name_col:
            return None

        name = row.get(name_col, "").strip()
        if not name:
            return None

        # Get optional fields
        description = None
        if column_mapping.get("description"):
            desc = row.get(column_mapping["description"], "").strip()
            description = desc if desc else None

        source_platform = None
        if column_mapping.get("source_platform"):
            sp = row.get(column_mapping["source_platform"], "").strip()
            source_platform = sp if sp else None

        install_status = None
        if column_mapping.get("install_status"):
            ist = row.get(column_mapping["install_status"], "").strip()
            install_status = ist if ist else None

        return UseCase(
            project_id=project_id,
            name=name,
            description=description,
            source_platform=source_platform,
            install_status=install_status,
        )


# Singleton instance
file_storage_service = FileStorageService()
